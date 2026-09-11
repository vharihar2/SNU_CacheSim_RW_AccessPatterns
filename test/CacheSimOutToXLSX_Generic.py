import sys
import os
import re
import argparse
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def parse_config(config_path):
    """
    Parses the configuration file to extract the number of cache levels
    and their respective capacities.
    """
    config = {}
    with open(config_path, 'r') as f:
        for line in f:
            # Remove inline comments and whitespace
            line = line.split('#')[0].strip()
            if not line or '=' not in line:
                continue
            key, val = line.split('=', 1)
            config[key.strip()] = val.strip()

    num_levels = int(config.get('NumCacheLevelsExclMM', 0))
    capacities = []
    
    for lvl in range(1, num_levels + 1):
        cap_key = f'CapacityOfL{lvl}Cache'
        if cap_key in config:
            capacities.append((lvl, int(config[cap_key])))
        else:
            raise ValueError(f"Missing required key '{cap_key}' in configuration file: {config_path}")

    return capacities

def parse_and_export_cache_log(file_path, config_path, output_excel):
    capacities = parse_config(config_path)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cache State"

    # Define Styles
    font_bold_red = Font(name="Calibri", size=11, bold=True, color="FF0000") # D:T, MCA:T
    font_bold = Font(name="Calibri", size=11, bold=True, color="000000")     # MCA:T only
    font_red = Font(name="Calibri", size=11, color="FF0000")              # D:T only
    font_normal = Font(name="Calibri", size=11, color="555555")              # Default
    fill_invalid = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid") # V:F
    
    # Header Styles
    font_header = Font(name="Calibri", size=11, bold=True, underline="single", color="000000")
    fill_header = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid") # Light Yellow

    # Regex to match individual cache entry tokens
    entry_pattern = re.compile(r'(-?\d+,-?\d+)\s*\(V:([TF]),\s*D:([TF]),\s*MCA:([TF]),\s*TS:(-?\d+),\s*LAT:(-?\d+)\)')

    # Dynamically build header list based on parsed capacities
    headers = ["Serial #"]
    for lvl, cap in capacities:
        for entry_idx in range(1, cap + 1):
            headers.append(f"L{lvl} Entry {entry_idx}")

    with open(file_path, 'r') as f:
        lines = f.readlines()

    ws.append(headers)

    # Format Header Row Cells
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_idx, line in enumerate(lines, start=2):
        if not line.strip():
            continue
        
        # Extract Serial Number
        serial_match = re.match(r'^\s*(\d+)\.', line)
        serial_num = serial_match.group(1) if serial_match else str(row_idx - 1)
        
        ws.cell(row=row_idx, column=1, value=int(serial_num)).alignment = Alignment(horizontal="center")
        
        # Find all cache entries in order
        entries = entry_pattern.findall(line)
        
        for col_idx, (addr_data, valid, dirty, mca, ts, lat) in enumerate(entries, start=2):
            cell = ws.cell(row=row_idx, column=col_idx)
            
            # Formats address pair and all attributes including TS and LAT
            cell.value = f"{addr_data}\n(V:{valid}, D:{dirty}, MCA:{mca}, TS:{ts}, LAT:{lat})"
            cell.alignment = Alignment(wrap_text=True, horizontal="center")
            
            # Formatting Rules
            if dirty == 'T':
                if mca == 'T':
                    cell.font = font_bold_red
                else:
                    cell.font = font_red
            elif dirty == 'F':
                if mca == 'T':
                    cell.font = font_bold
                else:
                    cell.font = font_normal
            else:
                cell.font = font_normal
                
            if valid == 'F':
                cell.fill = fill_invalid

    # Set Column Widths (7.78 units for Serial #, 27.78 units for all cache entries)
    ws.column_dimensions['A'].width = 7.78
    for col in range(2, ws.max_column + 1):
        col_letter = get_column_letter(col)
        ws.column_dimensions[col_letter].width = 27.78

    # Set Row Heights (28.8 pt for all rows including header)
    for row in range(1, ws.max_row + 1):
        ws.row_dimensions[row].height = 28.8

    wb.save(output_excel)
    print(f"Successfully generated: {output_excel}")

def main():
    parser = argparse.ArgumentParser(description="Visualize Cache Log file into a styled Excel sheet based on config settings.")
    parser.add_argument("input_file", help="Path to the input text log file")
    parser.add_argument("-c", "--config", help="Path to the config file (default: config_file.dat)", default="config_file.dat")
    parser.add_argument("-o", "--output", help="Optional output XLSX file path", default=None)

    args = parser.parse_args()

    if not os.path.exists(args.input_file):
        print(f"Error: Input file '{args.input_file}' does not exist.")
        sys.exit(1)

    if not os.path.exists(args.config):
        print(f"Error: Config file '{args.config}' does not exist.")
        sys.exit(1)

    output_file = args.output
    if not output_file:
        base_name = os.path.splitext(args.input_file)[0]
        output_file = f"{base_name}_visualized.xlsx"

    parse_and_export_cache_log(args.input_file, args.config, output_file)

if __name__ == "__main__":
    main()